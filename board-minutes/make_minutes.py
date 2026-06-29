# -*- coding: utf-8 -*-
import os
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

FONT = "Apple SD Gothic Neo"  # fallback to system default if missing

doc = Document()

# --- A4 paper ---
section = doc.sections[0]
section.page_height = Mm(297)
section.page_width = Mm(210)
section.top_margin = Mm(25)
section.bottom_margin = Mm(25)
section.left_margin = Mm(25)
section.right_margin = Mm(25)

# --- base style / font (incl. East Asian) ---
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

def set_font(run, size=11, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), FONT)
    rFonts.set(qn("w:hAnsi"), FONT)
    rFonts.set(qn("w:eastAsia"), FONT)

def para(text="", align=None, bold=False, size=11, space_after=6, space_before=0):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = 1.5
    if text:
        r = p.add_run(text)
        set_font(r, size=size, bold=bold)
    return p

CENTER = WD_ALIGN_PARAGRAPH.CENTER

# --- Title ---
para("이사회 의사록", align=CENTER, bold=True, size=18, space_after=18)

# --- Company ---
para("주식회사 원윌앤코", align=CENTER, size=12, space_after=18)

# --- Items 1-3 ---
para("1. 일    시 : 2026년 6월 1일      09시  30분")
para("2. 장    소 : 본점 (경기도 하남시 위례학암로 14번길 27, 5층 501호)")
para("3. 이사 총수 : 3명     출석이사 : 3명 (전원 출석)", space_after=14)

# --- Opening declaration ---
para("의장인 대표이사 박다빈은 이사회가 적법하게 성립되었음을 고하고 개회를 선언한 후, 다음 의안을 부의하다.", space_after=14)

# --- Agenda item 1 (bold) ---
para("제1호 의안 : 지점 설치의 건", bold=True, space_after=6)

p = para(space_after=10)
p.paragraph_format.left_indent = Mm(5)
r = p.add_run("의장은 회사의 영업상 필요에 따라 다음과 같이 지점을 설치할 것을 제안하였고 출석이사 전원의 찬성으로 이를 가결하다.")
set_font(r)

p = para("설치할 지점 : 서울특별시 강남구 선릉로86길 42, 3층(대치동)")
p.paragraph_format.left_indent = Mm(5)
p = para("지점 설치일 : 2026년 6월 15일", space_after=14)
p.paragraph_format.left_indent = Mm(5)

# --- Closing ---
para("이상으로 의안을 모두 의결하였으므로 의장은 폐회를 선언하다.", space_after=24)

# --- Date (center) ---
para("2026년 6월 1일", align=CENTER, space_after=18)

# --- Board (center) ---
para("주식회사 원윌앤코 이사회", align=CENTER, space_after=24)

# --- Signatures (인 stays blank for physical seal) ---
p = para("의 장  대표이사   박 다 빈        (인)")
p.paragraph_format.left_indent = Mm(5)
p = para("이    사            지 윤 진        (인)")
p.paragraph_format.left_indent = Mm(5)
p = para("이    사            송 채 훈        (인)")
p.paragraph_format.left_indent = Mm(5)

out = os.path.expanduser("~/Desktop/이사회의사록_지점설치.docx")
doc.save(out)
print("SAVED:", out)
