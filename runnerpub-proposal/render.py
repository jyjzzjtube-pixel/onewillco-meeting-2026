from pptx import Presentation
from pptx.util import Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image, ImageDraw, ImageFont
import io, os

EMU=914400.0
SC=155           # px per inch
W,H=int(13.333*SC), int(7.5*SC)
FC={}
def fnt(pt,bold):
    k=(round(pt*2),bold)
    if k not in FC:
        p='/root/.fonts/Pretendard-%s.otf'%('Bold' if bold else 'Regular')
        FC[k]=ImageFont.truetype(p,max(7,int(round(pt*SC/72))))
    return FC[k]
def rgb(c):
    try:
        if c and c.type is not None and c.rgb is not None: return '#'+str(c.rgb)
    except Exception: pass
    return None

def wrap(dr,txt,f,maxw):
    out=[];cur=''
    for ch in txt:
        if ch=='\n':
            out.append(cur);cur='';continue
        if dr.textlength(cur+ch,font=f)<=maxw: cur+=ch
        else: out.append(cur);cur=ch
    out.append(cur);return out

import sys
SRC=sys.argv[1] if len(sys.argv)>1 else 'runnerpub_v6.pptx'
OUTD=sys.argv[2] if len(sys.argv)>2 else '.'
os.makedirs(OUTD,exist_ok=True)
prs=Presentation(SRC)
for i,slide in enumerate(prs.slides,1):
    img=Image.new('RGB',(W,H),'white')
    if False:
        bgf='bg_cover.jpg' if i==1 else 'bg_end.jpg'
        img=Image.open(bgf).convert('RGB').resize((W,H))
    dr=ImageDraw.Draw(img,'RGBA')
    # 배경
    try:
        b=slide.background.fill
        if b.type is not None and rgb(b.fore_color): dr.rectangle([0,0,W,H],fill=rgb(b.fore_color))
    except Exception: pass
    for sh in slide.shapes:
        if sh.shape_type==13 or sh.__class__.__name__=='Picture':
            try:
                im=Image.open(io.BytesIO(sh.image.blob)).convert('RGBA')
                bx=(int(sh.left/EMU*SC),int(sh.top/EMU*SC),int(sh.width/EMU*SC),int(sh.height/EMU*SC))
                if bx[2]>W*0.9 and bx[3]>H*0.9: img.paste(im.resize((W,H)).convert('RGB'),(0,0))
                else:
                    im=im.resize((max(1,bx[2]),max(1,bx[3])))
                    img.paste(im,(bx[0],bx[1]),im)
            except Exception: pass
            continue
        if sh.left is None: continue
        if sh.has_table:
            tx,ty=sh.left/EMU*SC, sh.top/EMU*SC
            tb=sh.table
            colw=[c.width/EMU*SC for c in tb.columns]
            rowh=[r.height/EMU*SC for r in tb.rows]
            cy0=ty
            for ri,row in enumerate(tb.rows):
                cx0=tx
                for ci,cell in enumerate(row.cells):
                    fillc=None
                    try:
                        if cell.fill.type==1 and rgb(cell.fill.fore_color): fillc=rgb(cell.fill.fore_color)
                    except Exception: pass
                    if fillc: dr.rectangle([cx0,cy0,cx0+colw[ci],cy0+rowh[ri]],fill=fillc)
                    dr.rectangle([cx0,cy0,cx0+colw[ci],cy0+rowh[ri]],outline='#E3DFF2',width=1)
                    para=cell.text_frame.paragraphs[0]
                    rs=[r for r in para.runs if r.text]
                    if rs:
                        szz=max([(r.font.size.pt if r.font.size else 9.6) for r in rs])
                        bb=any(bool(r.font.bold) for r in rs)
                        cc=None
                        for r in rs:
                            c2=rgb(r.font.color)
                            if c2: cc=c2;break
                        ff=fnt(szz,bb); txt2=''.join(r.text for r in rs)
                        for k,l2 in enumerate(wrap(dr,txt2,ff,colw[ci]-16)):
                            dr.text((cx0+8,cy0+6+k*szz*SC/72*1.3),l2,font=ff,fill=cc or '#3A3550')
                    cx0+=colw[ci]
                cy0+=rowh[ri]
            continue
        x,y=sh.left/EMU*SC, sh.top/EMU*SC
        w,h=(sh.width or 0)/EMU*SC,(sh.height or 0)/EMU*SC
        # 도형 채우기
        try:
            f=sh.fill
            col=rgb(f.fore_color) if f.type==1 else None
            if col:
                st=str(sh.shape_type)
                alpha=255
                try:
                    xml=sh._element.xml
                    if 'alpha val=' in xml:
                        import re as _re
                        m=_re.search(r'alpha val="(\d+)"',xml)
                        if m: alpha=int(int(m.group(1))/100000*255)
                except Exception: pass
                fillrgba=tuple(int(col[j:j+2],16) for j in (1,3,5))+(alpha,)
                if 'OVAL' in st or 'ELLIPSE' in st: dr.ellipse([x,y,x+w,y+h],fill=fillrgba)
                else: dr.rounded_rectangle([x,y,x+w,y+h],radius=6,fill=fillrgba)
            ln=sh.line
            if ln.fill.type==1 and rgb(ln.color):
                dr.rounded_rectangle([x,y,x+w,y+h],radius=6,outline=rgb(ln.color),width=1)
        except Exception: pass
        # 텍스트
        if not sh.has_text_frame: continue
        tf=sh.text_frame; cy=y+2
        paras=[p for p in tf.paragraphs]
        # valign 중앙 처리용 총높이 추정
        total=0; plans=[]
        for p in paras:
            runs=[r for r in p.runs if r.text!='']
            if not runs: total+=6; plans.append(None); continue
            sz=max([(r.font.size.pt if r.font.size else 12) for r in runs])
            bold=any(bool(r.font.bold) for r in runs)
            txt=''.join(r.text for r in runs)
            f0=fnt(sz,bold)
            lines=wrap(dr,txt,f0,max(10,w-4))
            lh=sz*SC/72*1.35
            plans.append((lines,f0,lh,runs,sz))
            total+=lh*len(lines)
        try: anc=tf.vertical_anchor
        except Exception: anc=None
        if anc==MSO_ANCHOR.MIDDLE: cy=y+(h-total)/2
        for pl,p in zip(plans,paras):
            if pl is None: cy+=6; continue
            lines,f0,lh,runs,sz=pl
            col=None
            for r in runs:
                c=rgb(r.font.color)
                if c: col=c;break
            col=col or '#333333'
            al=p.alignment
            for ln_ in lines:
                tw=dr.textlength(ln_,font=f0)
                tx=x+2
                if al==PP_ALIGN.CENTER: tx=x+(w-tw)/2
                elif al==PP_ALIGN.RIGHT: tx=x+w-tw-2
                dr.text((tx,cy),ln_,font=f0,fill=col)
                cy+=lh
    img.save(os.path.join(OUTD,f'p{i}.png'))
print('rendered', len(prs.slides))
